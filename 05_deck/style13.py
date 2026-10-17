import copy, re
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from lxml import etree

NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main', 'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'p14': 'http://schemas.microsoft.com/office/powerpoint/2010/main'}
A = '{%s}' % NS['a']
IC = '/home/claude/v13/'
NAVY, DARK, GRAY, LG, LINE = RGBColor(0x0B, 0x4E, 0xA2), RGBColor(0x40, 0x40, 0x40), RGBColor(0x7F, 0x8C, 0x99), 'F2F2F2', RGBColor(0xCF, 0xD8, 0xE1)

EMPH = {
 2: ["높은 종 분리 성능", "녹음 비독립성", "녹음 조건 편향"],
 3: ["분석 대상 10,694클립", "표본 3,539클립", "교차검증 단위", "녹음 출처"],
 4: ["Linear probe", "Baseline 대비", "특징 분류기 성능"],
 5: ["레퍼런스 모델", "MFCC", "물리 특성"],
 6: ["8개 콜 타입", "3개 신호 유형"],
 7: ["녹음 출처", "종 단위", "샘플레이트", "분류군 단위", "녹음 조건"],
 8: ["랜덤 분할", "선행연구 수준이 재현", "출처를 분리", "baseline 수준"],
 9: ["학습 출처에 종속", "일반화되지 않음"],
 10: ["샘플레이트를 통제", "신호 유형 3분류", "baseline을 유의하게 상회"],
 11: ["EfficientNet-B3·AST", "baseline을 넘고", "growl·pulsed call"],
 12: ["단일 물리 특징", "대역", "1–8 kHz에서 중첩"],
 13: ["물리 특징값", "신호 유형", "분류군 확률", "주석 절차"],
 14: ["녹음 식별", "다출처 자료"],
 15: ["1,212개 패치", "모든 패치 사이의 관계", "12층"],
 16: ["0.96초 패치", "국소 패턴", "128차원 벡터"],
 17: ["0–20 kHz 스펙트로그램", "사진처럼", "26개 블록"],
 18: ["학습 없이", "20개 계수", "시간 통계 80개"],
 19: ["물리 특징 51개", "원 샘플레이트", "초·Hz 단위", "6개 그룹"],
 20: ["재현한 부분", "재현하지 않은 부분"],
 21: ["랜덤 분할", "재현", "출처 단위 분할", "chance 수준"],
 22: ["녹음 조건만으로도", "BA 0.38", "샘플레이트"],
 23: ["3유형에서 baseline을 넘지만(0.61)", "유의한 차이가 없음(0.46)"],
 24: ["BA 0.34", "baseline보다 낮음"],
 25: ["BA 0.37", "baseline 수준", "재현율이 낮음"],
 26: ["가장 높음", "BA 0.55", "whistle 재현율 0.81"],
 27: ["BA 0.48", "baseline을 넘고", "재현율이 가장 높음"],
 28: ["0–4 kHz", "바뀌지 않음"],
 29: ["거의 겹치지 않음", "20 kHz뿐"],
 30: ["녹음 출처 의존적", "신호 유형·대역 수준"],
}
SEC = [("재료 및 방법", "mm"), ("배경", "bg"), ("결과", "res"), ("고찰", "dis"), ("결론", "con"), ("부록", "app"), ("연구 요약", "sum"), ("목차", "toc")]

def color_el(rPr, blue):
    for f in rPr.findall(A + 'solidFill'): rPr.remove(f)
    sf = etree.SubElement(rPr, A + 'solidFill')
    if blue: etree.SubElement(sf, A + 'schemeClr', val='accent1')
    else:
        c = etree.SubElement(sf, A + 'schemeClr', val='tx1'); etree.SubElement(c, A + 'lumMod', val='90000'); etree.SubElement(c, A + 'lumOff', val='10000')
    # solidFill must come first among rPr children (before latin etc.)
    rPr.remove(sf); rPr.insert(0, sf)

def spans(text, phrases):
    mask = [False] * len(text)
    for ph in phrases:
        i = text.find(ph); assert i >= 0, (ph, text)
        for k in range(i, i + len(ph)): mask[k] = True
    out = []; cur = None
    for ch, m in zip(text, mask):
        if cur and cur[1] == m: cur[0] += ch
        else: cur = [ch, m]; out.append(cur)
    return out

def restyle_title(sh, phrases):
    txb = sh.text_frame._txBody; p = txb.find(A + 'p'); runs = p.findall(A + 'r')
    text = ''.join(r.find(A + 't').text or '' for r in runs)
    sizes = [r.find(A + 'rPr').get('sz') for r in runs if r.find(A + 'rPr') is not None and r.find(A + 'rPr').get('sz')]
    sz = int(sizes[0]) if sizes else 2200
    tmpl = copy.deepcopy(runs[0].find(A + 'rPr'))
    for r in runs: p.remove(r)
    end = p.find(A + 'endParaRPr')
    sz = max(1600, sz - 100)
    for seg, blue in spans(text, phrases):
        r = etree.Element(A + 'r'); rp = copy.deepcopy(tmpl); rp.set('sz', str(sz)); rp.set('b', '1')
        for k in ('err',):
            if k in rp.attrib: del rp.attrib[k]
        color_el(rp, blue); r.append(rp); t = etree.SubElement(r, A + 't'); t.text = seg
        if end is not None: end.addprevious(r)
        else: p.append(r)
    # light gray box with inset
    spPr = sh._element.spPr
    for f in spPr.findall(A + 'noFill') + spPr.findall(A + 'solidFill'): spPr.remove(f)
    sf = etree.Element(A + 'solidFill'); etree.SubElement(sf, A + 'srgbClr', val=LG)
    pg = spPr.find(A + 'prstGeom')
    if pg is None:
        pg = etree.SubElement(spPr, A + 'prstGeom', prst='rect'); etree.SubElement(pg, A + 'avLst')
    pg.addnext(sf)
    bp = txb.find(A + 'bodyPr'); bp.set('lIns', str(Inches(0.12))); bp.set('rIns', str(Inches(0.08)))
    return text

def restyle_tag(slide, sh):
    text = sh.text_frame.text; key = next((k for s, k in SEC if text.startswith(s)), None)
    m = re.match(r'^(부록 A\d+|연구 요약|결론 및 제언|목차|재료 및 방법(?: [①-⑨])?|\S+(?: [①-⑨])?)\s*(.*)$', text)
    head, sub = (m.group(1), m.group(2)) if m else (text, '')
    tf = sh.text_frame; p = tf.paragraphs[0]
    for r in list(p.runs): p._p.remove(r._r)
    r1 = p.add_run(); r1.text = head; r1.font.size = Pt(16); r1.font.bold = True; r1.font.color.rgb = NAVY
    if sub:
        r2 = p.add_run(); r2.text = '   ' + sub; r2.font.size = Pt(14); r2.font.bold = True; r2.font.color.rgb = DARK
    sh.left, sh.top, sh.width, sh.height = Inches(0.6 + 0.45), Inches(0.17), Inches(6.6), Inches(0.38)
    tf._txBody.find(A + 'bodyPr').set('anchor', 'ctr')
    if key: slide.shapes.add_picture(IC + f'ic_{key}.png', Inches(0.6), Inches(0.18), Inches(0.36), Inches(0.36))

def top_shapes(s):
    ph = [sh for sh in s.shapes if sh.is_placeholder and sh.has_text_frame]
    tag = next((sh for sh in ph if sh.placeholder_format.idx == 100), None)
    title = next((sh for sh in ph if sh.placeholder_format.idx == 101), None)
    return tag, title

prs = Presentation('/home/claude/v13/in.pptx')
for i, s in enumerate(prs.slides, 1):
    if i == 1: continue
    tag, title = top_shapes(s)
    if title is not None and i in EMPH: restyle_title(title, EMPH[i])
    if tag is not None: restyle_tag(s, tag)

# ---------- TOC slide ----------
lay = prs.slides[2].slide_layout
s = prs.slides.add_slide(lay)
tag, title = top_shapes(s)
tag.text_frame.text = '목차'; title.text_frame.text = '발표 구성'
title.left, title.top, title.width, title.height = Inches(0.6), Inches(0.68), Inches(12.13), Inches(0.55)
for r in title.text_frame.paragraphs[0].runs: r.font.size = Pt(22); r.font.bold = True
restyle_title(title, ['발표 구성']); restyle_tag(s, tag)
for sh in list(s.placeholders):
    if sh.placeholder_format.idx not in (100, 101): sh._element.getparent().remove(sh._element)
COLS = [("01", "배경", "bg", "p. 3", ["선행연구와 문제 정의"]),
        ("02", "재료 및 방법", "mm", "p. 4–7", ["데이터 구성 · 녹음 출처", "분석 흐름 · 평가 설계", "특징 추출기", "콜 타입 라벨"]),
        ("03", "결과", "res", "p. 8–13", ["① 데이터 구조", "② 재현성과 일반화 (RQ1·2)", "③ UMAP 일반화 (RQ2)", "④ 신호 유형 강건성 (RQ3)", "⑤ 콜 타입 6종", "⑥ 단일 특징과 대역"]),
        ("04", "고찰", "dis", "p. 14–15", ["주석 절차 제안", "결론 및 제언"]),
        ("+", "부록 · 요약", "app", "p. 16–31", ["A1–A5  특징 추출기 · 물리 특징", "A6–A7  선행연구 재현", "A8–A13  콜 타입 6종 상세", "A14  공통 대역 검증", "A15  분류군별 샘플레이트", "연구 요약 (p. 31)"])]
L0, W, TOP, BOT, G = 0.6, 12.13, 1.55, 6.75, 0.2
cw = (W - 4 * G) / 5
for k, (num, name, ic, pg, items) in enumerate(COLS):
    x = L0 + k * (cw + G)
    hd = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(TOP), Inches(cw), Inches(1.15)); hd.fill.solid(); hd.fill.fore_color.rgb = NAVY if k < 4 else RGBColor(0x5B, 0x67, 0x70); hd.line.fill.background()
    s.shapes.add_picture(IC + f'ic_{ic}.png', Inches(x + 0.15), Inches(TOP + 0.15), Inches(0.42), Inches(0.42))
    tb = s.shapes.add_textbox(Inches(x + 0.65), Inches(TOP + 0.08), Inches(cw - 0.7), Inches(0.55)); tf = tb.text_frame; tf.margin_left = tf.margin_right = 0
    r = tf.paragraphs[0].add_run(); r.text = num; r.font.size = Pt(26); r.font.bold = True; r.font.color.rgb = RGBColor(0xC8, 0xE9, 0xF8)
    tb2 = s.shapes.add_textbox(Inches(x + 0.15), Inches(TOP + 0.62), Inches(cw - 0.25), Inches(0.45)); tf2 = tb2.text_frame; tf2.margin_left = 0
    r = tf2.paragraphs[0].add_run(); r.text = name; r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    body = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(TOP + 1.15), Inches(cw), Inches(BOT - TOP - 1.15)); body.fill.solid(); body.fill.fore_color.rgb = RGBColor(0xF2, 0xF2, 0xF2); body.line.fill.background()
    tb3 = s.shapes.add_textbox(Inches(x + 0.15), Inches(TOP + 1.3), Inches(cw - 0.25), Inches(BOT - TOP - 1.9)); tf3 = tb3.text_frame; tf3.word_wrap = True; tf3.margin_left = 0
    for j, it in enumerate(items):
        p = tf3.paragraphs[0] if j == 0 else tf3.add_paragraph(); p.space_after = Pt(7)
        r = p.add_run(); r.text = it; r.font.size = Pt(12); r.font.color.rgb = DARK
    tb4 = s.shapes.add_textbox(Inches(x + 0.15), Inches(BOT - 0.45), Inches(cw - 0.25), Inches(0.35)); tf4 = tb4.text_frame; tf4.margin_left = 0
    r = tf4.paragraphs[0].add_run(); r.text = pg; r.font.size = Pt(11); r.font.bold = True; r.font.color.rgb = GRAY
# move TOC to position 2 and register in first section
sld = prs.slides._sldIdLst; ids = list(sld); new = ids[-1]; sld.remove(new); sld.insert(1, new)
nid = new.get('id')
ext = prs.part._element.find('.//p14:sectionLst', NS)
if ext is not None:
    secs = ext.findall('p14:section', NS)
    for sec in secs:
        lst = sec.find('p14:sldIdLst', NS)
        for e in list(lst):
            if e.get('id') == nid: lst.remove(e)
    first = secs[0].find('p14:sldIdLst', NS)
    etree.SubElement(first, '{%s}sldId' % NS['p14'], id=nid)
prs.save('/home/claude/v13/Watkins_lab_seminar_v13.pptx'); print('saved')
