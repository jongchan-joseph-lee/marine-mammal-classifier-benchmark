"""Helpers for editing the user's deck while matching its existing styles."""
import re
from copy import deepcopy
from lxml import etree
from pptx.util import Inches, Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

NAVY = "0B4EA2"; TXT = "2B2B2B"; GRAY = "7F8C99"; LINE = "CFD8E1"; SKY = "C8E9F8"; BAND = "EEF7FC"
ORANGE = "E07B00"; WHITE = "FFFFFF"; LIGHT = "F5F8FB"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"

def IN(v): return Inches(v)

# ---------------------------------------------------------------- slides
def dup_slide(prs, src):
    new = prs.slides.add_slide(src.slide_layout)
    for sh in list(new.shapes):
        sh._element.getparent().remove(sh._element)
    for el in src.shapes._spTree:
        if el.tag in (qn('p:nvGrpSpPr'), qn('p:grpSpPr')):
            continue
        new.shapes._spTree.append(deepcopy(el))
    for blip in new.shapes._spTree.iter(qn('a:blip')):
        rid = blip.get(qn('r:embed'))
        part = src.part.related_part(rid)
        blip.set(qn('r:embed'), new.part.relate_to(part, RT.IMAGE))
    return new

def slide_id(prs, slide):
    for sid in prs.slides._sldIdLst:
        if prs.part.related_part(sid.rId) is slide.part:
            return sid
def move_slide(prs, slide, idx):
    sl = prs.slides._sldIdLst
    sid = slide_id(prs, slide)
    sl.remove(sid); sl.insert(idx, sid)

def sections(prs):
    return prs.part._element.find('.//{%s}sectionLst' % P14)

def section_add(prs, name, slides, before_id=None, new_after=None):
    """Add slide ids to section `name` (before sldId before_id if given). If section is missing and
    new_after is given, create a new section after section `new_after`."""
    sl = sections(prs)
    sec = [s for s in sl if s.get('name') == name]
    if not sec:
        ref = [s for s in sl if s.get('name') == new_after][0]
        sec = etree.SubElement(sl, '{%s}section' % P14); sec.set('name', name)
        sec.set('id', '{7A0C1E55-2B6D-4C3A-9E1F-0D5B8A3C6E21}')
        etree.SubElement(sec, '{%s}sldIdLst' % P14)
        sl.remove(sec); ref.addnext(sec)
    else:
        sec = sec[0]
    lst = sec.find('{%s}sldIdLst' % P14)
    ids = [slide_id(prs, s).get('id') for s in slides]
    anchor = None
    if before_id is not None:
        anchor = [e for e in lst if e.get('id') == str(before_id)][0]
    for i in ids:
        e = etree.Element('{%s}sldId' % P14); e.set('id', i)
        if anchor is not None: anchor.addprevious(e)
        else: lst.append(e)

# ---------------------------------------------------------------- shapes
def rect(s, x, y, w, h, fill, line=None, lw=9525, geom=MSO_SHAPE.RECTANGLE, dash=None):
    sh = s.shapes.add_shape(geom, IN(x), IN(y), IN(w), IN(h))
    st = sh._element.find(qn('p:style'))
    if st is not None: st.getparent().remove(st)
    if fill is None: sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(fill)
    if line:
        sh.line.color.rgb = RGBColor.from_string(line); sh.line.width = Emu(lw)
        if dash:
            ln = sh._element.spPr.find(qn('a:ln'))
            pd = etree.SubElement(ln, qn('a:prstDash')); pd.set('val', dash)
    else:
        sh.line.fill.background()
    return sh

def _rpr(size, color, bold=False, scheme=None, italic=False):
    r = etree.Element(qn('a:rPr')); r.set('lang', 'ko-KR'); r.set('sz', str(int(round(size * 100))))
    if bold: r.set('b', '1')
    if italic: r.set('i', '1')
    r.set('dirty', '0')
    sf = etree.SubElement(r, qn('a:solidFill'))
    if scheme:
        sc = etree.SubElement(sf, qn('a:schemeClr')); sc.set('val', scheme[0])
        for k, v in scheme[1:]:
            e = etree.SubElement(sc, qn('a:' + k)); e.set('val', v)
    else:
        c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', color)
    return r

def _para(text, size, color, bold, align, emph, bullet, spc_after, line_sp=None, italic=False):
    p = etree.Element(qn('a:p'))
    ppr = etree.SubElement(p, qn('a:pPr'))
    if bullet:
        ppr.set('marL', '165100'); ppr.set('indent', '-165100')
    else:
        ppr.set('marL', '0'); ppr.set('indent', '0')
    ppr.set('algn', {'l': 'l', 'c': 'ctr', 'r': 'r'}[align])
    if line_sp:
        ls = etree.SubElement(ppr, qn('a:lnSpc')); sp = etree.SubElement(ls, qn('a:spcPct')); sp.set('val', str(int(line_sp * 1000)))
    if spc_after:
        sa = etree.SubElement(ppr, qn('a:spcAft')); sp = etree.SubElement(sa, qn('a:spcPts')); sp.set('val', str(int(spc_after * 100)))
    if bullet:
        b = etree.SubElement(ppr, qn('a:buSzPct')); b.set('val', '100000')
        b = etree.SubElement(ppr, qn('a:buChar')); b.set('char', '•')
    else:
        etree.SubElement(ppr, qn('a:buNone'))
    for k, seg in enumerate(re.split(r'\*\*', text)):
        if seg == '': continue
        is_emph = (k % 2 == 1)
        r = etree.SubElement(p, qn('a:r'))
        if is_emph:
            r.append(_rpr(size, emph, True, italic=italic))
        else:
            r.append(_rpr(size, color, bold, italic=italic))
        t = etree.SubElement(r, qn('a:t')); t.text = seg
    e = etree.SubElement(p, qn('a:endParaRPr')); e.set('lang', 'ko-KR'); e.set('sz', str(int(size * 100))); e.set('dirty', '0')
    return p

def fill_tf(sh, paras, size=11, color=TXT, bold=False, align='l', anchor='ctr', emph=NAVY, bullet=False,
            spc_after=0, line_sp=None, italic=False):
    if isinstance(paras, str): paras = [paras]
    txb = sh._element.find(qn('p:txBody'))
    bp = txb.find(qn('a:bodyPr'))
    for k in ('lIns', 'tIns', 'rIns', 'bIns'): bp.set(k, '0')
    bp.set('wrap', 'square'); bp.set('rtlCol', '0'); bp.set('anchor', {'t': 't', 'ctr': 'ctr', 'b': 'b'}[anchor])
    for c in list(bp): bp.remove(c)
    for p in txb.findall(qn('a:p')): txb.remove(p)
    for t in paras:
        txb.append(_para(t, size, color, bold, align, emph, bullet, spc_after, line_sp, italic))
    return sh

def tb(s, x, y, w, h, paras, **kw):
    sh = s.shapes.add_textbox(IN(x), IN(y), IN(w), IN(h))
    sh.text_frame.word_wrap = True
    return fill_tf(sh, paras, **kw)

def card(s, x, y, w, h, title, head=0.4, size=13, center=False):
    rect(s, x, y, w, head, NAVY, NAVY)
    rect(s, x, y + head, w, h - head, WHITE, LINE)
    tb(s, x + (0 if center else 0.15), y, w - (0 if center else 0.3), head, title, size=size, color=WHITE, bold=True,
       align='c' if center else 'l')

def picture(s, path, x, y, w=None, h=None):
    return s.shapes.add_picture(path, IN(x), IN(y), IN(w) if w else None, IN(h) if h else None)

# ---------------------------------------------------------------- text in existing template shapes
def set_title(sh, text):
    """[[...]] marks emphasis (accent1); the rest dark gray, as in the deck."""
    txb = sh._element.find(qn('p:txBody'))
    p = txb.find(qn('a:p'))
    for r in p.findall(qn('a:r')) + p.findall(qn('a:endParaRPr')): p.remove(r)
    for k, seg in enumerate(re.split(r'\[\[|\]\]', text)):
        if not seg: continue
        r = etree.SubElement(p, qn('a:r'))
        if k % 2 == 1: r.append(_rpr(16, None, True, scheme=('accent1',)))
        else: r.append(_rpr(16, None, True, scheme=('tx1', ('lumMod', '90000'), ('lumOff', '10000'))))
        t = etree.SubElement(r, qn('a:t')); t.text = seg
    e = etree.SubElement(p, qn('a:endParaRPr')); e.set('lang', 'ko-KR'); e.set('sz', '1600'); e.set('b', '1'); e.set('dirty', '0')

def set_tag(sh, head, sub):
    p = sh._element.find(qn('p:txBody')).find(qn('a:p'))
    runs = p.findall(qn('a:r'))
    runs[0].find(qn('a:t')).text = head
    if len(runs) < 2:
        r = etree.Element(qn('a:r')); r.append(_rpr(14, '404040', True))
        t = etree.SubElement(r, qn('a:t')); runs[0].addnext(r); runs.append(r)
    runs[1].find(qn('a:t')).text = '   ' + sub
    for r in runs[2:]: p.remove(r)

def set_plain(sh, text):
    p = sh._element.find(qn('p:txBody')).find(qn('a:p'))
    runs = p.findall(qn('a:r'))
    runs[0].find(qn('a:t')).text = text
    for r in runs[1:]: p.remove(r)

def set_band(sh, bullets):
    fill_tf(sh, bullets, size=13, color=TXT, bullet=True, spc_after=6, anchor='ctr')

def set_progress(sh, active=None):
    names = ['배경', '재료 및 방법', '결과', '고찰']
    txb = sh._element.find(qn('p:txBody')); p = txb.find(qn('a:p'))
    for r in p.findall(qn('a:r')): p.remove(r)
    end = p.find(qn('a:endParaRPr'))
    for i, n in enumerate(names):
        r = etree.Element(qn('a:r'))
        on = (n == active)
        rp = _rpr(10, NAVY if on else 'B8C2CC', on); r.append(rp)
        t = etree.SubElement(r, qn('a:t')); t.text = n + ('  ›  ' if i < 3 else '')
        end.addprevious(r)

NOTES_TEMPLATE = None
def notes(slide, text):
    ns = slide.notes_slide
    tf = ns.notes_text_frame
    if tf is None:
        ns.shapes._spTree.append(deepcopy(NOTES_TEMPLATE))
        tf = ns.notes_text_frame
    tf.text = text

# ---------------------------------------------------------------- table
def table(s, x, y, colw, rowh, rows, head_size=10.5, size=10, emph=NAVY, aligns=None, bold_col0=True, fills=None):
    nr, nc = len(rows), len(colw)
    gf = s.shapes.add_table(nr, nc, IN(x), IN(y), IN(sum(colw)), IN(sum(rowh)))
    tbl = gf._element.graphic.graphicData.tbl
    tp = tbl.tblPr
    for k in list(tp.attrib): del tp.attrib[k]
    for c in list(tp): tp.remove(c)
    for j, w in enumerate(colw): gf.table.columns[j].width = IN(w)
    for i, h in enumerate(rowh): gf.table.rows[i].height = IN(h)
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = gf.table.cell(i, j)
            al = 'c' if i == 0 else ((aligns or ['l'] * nc)[j])
            paras = val.split('\n')
            txb = cell._tc.txBody
            for p in txb.findall(qn('a:p')): txb.remove(p)
            for t in paras:
                if i == 0: txb.append(_para(t, head_size, WHITE, True, 'c', WHITE, False, 0))
                else:
                    sub = t.startswith('~')
                    tt = t[1:] if sub else t
                    txb.append(_para(tt, size - 1.5 if sub else size, GRAY if sub else TXT,
                                     (j == 0 and bold_col0 and not sub), al, emph, False, 0))
            tcPr = cell._tc.get_or_add_tcPr()
            for k, v in (('marL', '64008'), ('marR', '45720'), ('marT', '27432'), ('marB', '27432'), ('anchor', 'ctr')):
                tcPr.set(k, v)
            for c in list(tcPr): tcPr.remove(c)
            for side in ('lnL', 'lnR', 'lnT', 'lnB'):
                ln = etree.SubElement(tcPr, qn('a:' + side)); ln.set('w', '9525')
                sf = etree.SubElement(ln, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr')); c.set('val', LINE)
            sf = etree.SubElement(tcPr, qn('a:solidFill')); c = etree.SubElement(sf, qn('a:srgbClr'))
            if i == 0: c.set('val', NAVY)
            elif fills and fills.get((i, j)): c.set('val', fills[(i, j)])
            else: c.set('val', WHITE if i % 2 else LIGHT)
    return gf
