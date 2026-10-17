from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
NAVY=RGBColor(0x0B,0x4E,0xA2); DARK=RGBColor(0x40,0x40,0x40); GRAY=RGBColor(0x8A,0x94,0x9E); WHITE=RGBColor(255,255,255)
prs=Presentation('Watkins_lab_seminar_v13.pptx')
# keep only slide 2 (TOC), clear it
keep=prs.slides[1]
sl=prs.slides._sldIdLst
for sid in list(sl):
    if sid.get('id')!=sl[1].get('id'): pass
ids=list(sl); keep_id=ids[1]
for sid in ids:
    if sid is not keep_id:
        prs.part.drop_rel(sid.rId); sl.remove(sid)
ext=prs.part._element.find('.//{http://schemas.microsoft.com/office/powerpoint/2010/main}sectionLst')
if ext is not None: ext.getparent().getparent().remove(ext.getparent())
for sh in list(keep.shapes): sh._element.getparent().remove(sh._element)
s=keep
def box(x,y,w,h,text,size,color,bold=False,align=PP_ALIGN.LEFT,font=None):
    tb=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=tb.text_frame; tf.word_wrap=True; from pptx.enum.text import MSO_AUTO_SIZE; tf.auto_size=MSO_AUTO_SIZE.NONE; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0; tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    p=tf.paragraphs[0]; p.alignment=align; r=p.add_run(); r.text=text; r.font.size=Pt(size); r.font.bold=bold; r.font.color.rgb=color
    if font: r.font.name=font
    return tb
pan=s.shapes.add_shape(MSO_SHAPE.RECTANGLE,0,0,Inches(4.75),prs.slide_height); pan.fill.solid(); pan.fill.fore_color.rgb=NAVY; pan.line.fill.background(); pan.shadow.inherit=False
box(0.95,0.75,3.6,0.9,'Contents',44,WHITE,True,font='Arial Black')
box(0.98,1.6,3.4,0.4,'목차',16,RGBColor(0xC8,0xE9,0xF8),True)
COLS=[[("01","배경",[("선행연구와 문제 정의","03")]),
       ("02","재료 및 방법",[("데이터 구성 · 녹음 출처","04"),("분석 흐름 · 평가 설계","05"),("특징 추출기","06"),("콜 타입 라벨","07")]),
       ("+","부록 · 요약",[("특징 추출기 구조 (A1–5)","16"),("선행연구 재현 (A6–7)","21"),("콜 타입 6종 상세 (A8–13)","23"),("공통 대역 · 샘플레이트 (A14–15)","29"),("연구 요약","31")])],
      [("03","결과",[("데이터 구조","08"),("재현성과 일반화 (RQ1·RQ2)","09"),("UMAP 일반화 (RQ2)","10"),("신호 유형 강건성 (RQ3)","11"),("콜 타입 6종","12"),("단일 특징과 대역","13")]),
       ("04","고찰",[("주석 절차 제안","14"),("결론 및 제언","15")])]]
X0=[5.35,9.25]; HH,IH,GAP=0.5,0.34,0.24
for c,col in enumerate(COLS):
    y=1.05; x=X0[c]
    for num,name,items in col:
        box(x,y,0.55,HH,num,20,NAVY,True); box(x+0.6,y,3.0,HH,name,20,NAVY,True); y+=HH+0.04
        for it,pg in items:
            box(x+0.3,y,2.95,IH,'•  '+it,12.5,DARK); box(x+3.2,y,0.6,IH,'-- '+pg,12.5,GRAY,align=PP_ALIGN.RIGHT); y+=IH
        y+=GAP
prs.save('Watkins_TOC_v13.pptx'); print(len(prs.slides))
